import contextlib
import io
import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document

from metropole.chat import Chat
from metropole.cli import main
from metropole.config import Settings, ollama_url
from metropole.ingest import Procedure, chunk_text, parse_notes, read_source
from metropole.ollama import Ollama, OllamaError
from metropole.store import Store
from metropole.web import create_app


class StubOllama(BaseHTTPRequestHandler):
    """Local protocol fixture; does not claim to validate a real model."""
    calls = []
    release_stream = threading.Event()

    def log_message(self, *args):
        pass

    def respond(self, value):
        payload = json.dumps(value).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def do_GET(self):
        if self.path != "/api/tags":
            self.send_error(404)
            return
        self.respond({"models": [{"name": "test-model:latest"}]})

    def do_POST(self):
        if self.path != "/api/chat":
            self.send_error(404)
            return
        payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        self.calls.append(payload)
        naming = "Name a datacenter procedure" in payload["messages"][0]["content"]
        general = "using your general knowledge" in payload["messages"][0]["content"]
        if naming:
            answer = "rack-access-checklist"
        elif general:
            answer = "Choose a VM backup tool compatible with your hypervisor and test a restore."
        else:
            answer = "Confirm the backup date before restoring files. [1]"
        if payload.get("stream"):
            events = [
                {"message": {"content": answer[:20]}, "done": False},
                {"message": {"content": answer[20:]}, "done": False},
                {"message": {"content": ""}, "done": True},
            ]
            body = "".join(json.dumps(event) + "\n" for event in events).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/x-ndjson")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            first, remaining = body.split(b"\n", 1)
            self.wfile.write(first + b"\n")
            self.wfile.flush()
            if "PAUSE_STREAM" in payload["messages"][-1]["content"]:
                self.release_stream.wait(timeout=5)
            self.wfile.write(remaining)
        else:
            self.respond({"message": {"role": "assistant", "content": answer}})


class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), StubOllama)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.settings = Settings(host=f"127.0.0.1:{self.server.server_port}", database=str(self.root / "kb.sqlite3"))
        self.store = Store(self.settings.db_path)
        self.client = Ollama(self.settings)
        StubOllama.calls.clear()

    def tearDown(self):
        self.temp.cleanup()

    def seed(self):
        self.store.replace_source("notes.txt", [Procedure("restore-files-from-avamar", "Confirm the backup date before restoring files with Avamar.", "notes.txt", 1)])

    def test_host_configuration(self):
        self.assertEqual(ollama_url("10.3.81.142"), "http://10.3.81.142:11434")
        self.assertEqual(ollama_url("https://ollama.example:443"), "https://ollama.example:443")
        for host in ["", "ftp://server", "http://a/b", "http://user:pass@server"]:
            with self.assertRaises(ValueError):
                ollama_url(host)

    def test_mixed_notes_preserve_order_and_infer_tag(self):
        text = "<restore-files-from-avamar>Restore text</restore-files-from-avamar>\n------\nCheck rack access"
        procedures = parse_notes(text, "notes", self.client)
        self.assertEqual([p.tag for p in procedures], ["restore-files-from-avamar", "rack-access-checklist"])
        self.assertEqual(procedures[1].text, "Check rack access")
        self.assertEqual(len(StubOllama.calls), 1)

    def test_unmatched_text_before_and_after_tags_is_not_lost(self):
        items = parse_notes("Before\n<a>Tagged</a>\nAfter", "notes", fallback=True)
        self.assertEqual([p.text for p in items], ["Before", "Tagged", "After"])

    def test_utf8_bom_and_offline_tags(self):
        path = self.root / "notes.txt"
        path.write_text("Check rack access\nConfirm the access window", encoding="utf-8-sig")
        self.assertEqual(read_source(path, fallback=True)[0].tag, "check-rack-access")

    def test_tagged_notes_need_no_model(self):
        items = parse_notes("<restore-files>Do a restore</restore-files>", "notes")
        self.assertEqual(items[0].tag, "restore-files")
        self.assertEqual(StubOllama.calls, [])

    def test_unlabelled_notes_without_client_fail(self):
        with self.assertRaises(ValueError):
            parse_notes("A procedure", "notes")

    def test_docx_filename_paragraphs_and_tables(self):
        path = self.root / "Restore Files From Avamar.docx"
        document = Document()
        document.add_paragraph("Before table")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Client"
        table.cell(0, 1).text = "Backup date"
        document.add_paragraph("After table")
        document.save(path)
        item = read_source(path)[0]
        self.assertEqual(item.tag, "restore-files-from-avamar")
        self.assertEqual(item.text, "Before table\nClient | Backup date\nAfter table")

    def test_chunking_long_text_retains_tail(self):
        text = "restoration " * 1000 + "UNIQUE_TAIL"
        chunks = list(chunk_text(text))
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(c) <= 1800 for c in chunks))
        self.assertTrue(chunks[-1].endswith("UNIQUE_TAIL"))

    def test_reimport_replaces_stale_chunks(self):
        self.seed()
        self.seed()
        self.assertEqual(len(self.store.list_procedures()), 1)
        self.store.replace_source("notes.txt", [Procedure("rack-access", "Rack access requires approval", "notes.txt", 1)])
        self.assertEqual(self.store.search("Avamar"), [])
        self.assertEqual(len(self.store.search("rack access")), 1)

    def test_failed_import_rolls_back(self):
        self.seed()
        with self.assertRaises(Exception):
            self.store.replace_source("notes.txt", [Procedure("a", "text", "notes.txt", 1), Procedure("b", "text", "notes.txt", 1)])
        self.assertEqual(self.store.list_procedures()[0]["tag"], "restore-files-from-avamar")

    def test_search_treats_special_characters_as_data(self):
        self.seed()
        self.assertTrue(self.store.search('Avamar " OR * : ()'))
        self.assertEqual(self.store.search("???"), [])

    def test_grounded_chat_and_followup(self):
        self.seed()
        chat = Chat(self.store, self.client)
        result = chat.ask("How do I restore Avamar files?")
        self.assertIn("[1]", result["answer"])
        self.assertNotIn("General knowledge", result["answer"])
        self.assertIn("Use only the supplied excerpts", StubOllama.calls[-1]["messages"][0]["content"])
        self.assertEqual(result["sources"][0]["tag"], "restore-files-from-avamar")
        history = [{"role": "user", "content": "How do I restore Avamar files?"}, {"role": "assistant", "content": result["answer"]}]
        self.assertTrue(chat.ask("Which date?", history)["sources"])
        sent = StubOllama.calls[-1]
        self.assertEqual(sent["stream"], True)
        self.assertIn("Reference excerpts", sent["messages"][-1]["content"])
        self.assertEqual(sent["messages"][1], history[0])

    def test_no_match_uses_labelled_general_knowledge(self):
        self.seed()
        result = Chat(self.store, self.client).ask("Switch firmware upgrades")
        self.assertEqual(result["sources"], [])
        self.assertEqual(len(StubOllama.calls), 1)
        self.assertTrue(result["answer"].startswith("General knowledge (no matching internal procedure found):"))
        self.assertIn("Choose a VM backup tool", result["answer"])
        self.assertNotIn("[1]", result["answer"])
        sent = StubOllama.calls[0]
        self.assertIn("No matching procedure", sent["messages"][0]["content"])
        self.assertEqual(sent["messages"][-1]["content"], "Switch firmware upgrades")

    def test_empty_repository_general_knowledge_and_followup(self):
        chat = Chat(self.store, self.client)
        result = chat.ask("How to backup a VM?")
        history = [{"role": "user", "content": "How to backup a VM?"},
                   {"role": "assistant", "content": result["answer"]}]
        followup = chat.ask("What about testing it?", history)
        self.assertEqual(followup["sources"], [])
        self.assertEqual(StubOllama.calls[-1]["messages"][1:3], history)
        self.assertIn("General knowledge", followup["answer"])

    def test_flask_no_match_uses_general_knowledge(self):
        browser = create_app(self.settings).test_client()
        result = browser.post("/api/chat", json={"question": "How to backup a VM?"})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json["sources"], [])
        self.assertIn("General knowledge", result.json["answer"])

    def test_general_knowledge_model_failure_returns_error(self):
        with patch.object(Ollama, "chat", side_effect=OllamaError("server unavailable")):
            browser = create_app(self.settings).test_client()
            result = browser.post("/api/chat", json={"question": "How to backup a VM?"})
            self.assertEqual(result.status_code, 502)

    def test_explicit_missing_model_reports_error(self):
        client = Ollama(Settings(self.settings.host, "missing"))
        with self.assertRaises(OllamaError):
            client.choose_model()

    def test_ollama_yields_before_response_finishes(self):
        StubOllama.release_stream.clear()
        stream = self.client.chat_stream([{"role": "user", "content": "PAUSE_STREAM"}])
        try:
            with ThreadPoolExecutor(max_workers=1) as executor:
                future = executor.submit(next, stream)
                try:
                    first = future.result(timeout=2)
                    self.assertEqual(first, "Confirm the backup d")
                    self.assertFalse(StubOllama.release_stream.is_set())
                finally:
                    StubOllama.release_stream.set()
            self.assertEqual(first + "".join(stream), "Confirm the backup date before restoring files. [1]")
        finally:
            StubOllama.release_stream.set()
            stream.close()

    def test_flask_stream_events_and_general_knowledge(self):
        browser = create_app(self.settings).test_client()
        response = browser.post("/api/chat", json={"question": "Backup a VM?", "stream": True}, buffered=False)
        self.assertEqual(response.mimetype, "application/x-ndjson")
        events = [json.loads(line) for line in response.response]
        self.assertEqual([e["type"] for e in events], ["metadata", "token", "token", "done"])
        self.assertEqual(events[0]["sources"], [])
        self.assertIn("General knowledge", events[0]["prefix"])
        self.assertEqual("".join(e["text"] for e in events if e["type"] == "token"),
                         "Choose a VM backup tool compatible with your hypervisor and test a restore.")
        response.close()

    def test_flask_stream_preserves_internal_sources(self):
        self.seed()
        browser = create_app(self.settings).test_client()
        response = browser.post("/api/chat", json={"question": "Avamar restore", "stream": True})
        events = [json.loads(line) for line in response.data.splitlines()]
        self.assertEqual(events[0]["sources"][0]["tag"], "restore-files-from-avamar")
        self.assertEqual(events[0]["prefix"], "")
        self.assertEqual(events[-1]["type"], "done")

    def test_stream_errors_are_terminal_and_partial_text_is_preserved(self):
        def failing_stream(*args, **kwargs):
            yield "Partial answer"
            raise OllamaError("connection lost")
        with patch.object(Ollama, "chat_stream", side_effect=failing_stream):
            browser = create_app(self.settings).test_client()
            response = browser.post("/api/chat", json={"question": "Backup a VM", "stream": True})
            events = [json.loads(line) for line in response.data.splitlines()]
            self.assertEqual([e["type"] for e in events], ["metadata", "token", "error"])
            self.assertEqual(events[1]["text"], "Partial answer")
            self.assertEqual(events[-1]["error"], "connection lost")

    def test_stream_validation_happens_before_headers(self):
        browser = create_app(self.settings).test_client()
        for payload in [{"question": "", "stream": True}, {"question": "VM", "stream": "yes"}]:
            result = browser.post("/api/chat", json=payload)
            self.assertEqual(result.status_code, 400)
            self.assertEqual(result.mimetype, "application/json")

    def test_incomplete_and_malformed_ollama_streams_fail(self):
        self.client.model = "test-model:latest"
        for lines in [[b'{"message":{"content":"partial"}}'], [b'not json'],
                      [b'{"error":"model crashed"}'], [b'{"done":true}']]:
            response = MagicMock()
            response.__enter__.return_value = response
            response.iter_lines.return_value = iter(lines)
            with patch("metropole.ollama.requests.post", return_value=response):
                with self.assertRaises(OllamaError):
                    list(self.client.chat_stream([{"role": "user", "content": "VM backup"}]))
            response.__exit__.assert_called_once()

    def test_cancelling_stream_closes_upstream_connection(self):
        self.client.model = "test-model:latest"
        response = MagicMock()
        response.__enter__.return_value = response
        response.iter_lines.return_value = iter([b'{"message":{"content":"first"}}', b'{"done":true}'])
        with patch("metropole.ollama.requests.post", return_value=response):
            stream = self.client.chat_stream([{"role": "user", "content": "VM"}])
            self.assertEqual(next(stream), "first")
            stream.close()
        response.__exit__.assert_called_once()

    def test_flask_chat_status_and_validation(self):
        self.seed()
        browser = create_app(self.settings).test_client()
        self.assertEqual(browser.get("/").status_code, 200)
        self.assertEqual(browser.get("/health").json["procedures"], 1)
        self.assertEqual(browser.get("/api/status").status_code, 200)
        result = browser.post("/api/chat", json={"question": "Restore Avamar files"})
        self.assertEqual(result.status_code, 200)
        self.assertTrue(result.json["sources"])
        for data in [[], {}, {"question": ""}, {"question": "x", "history": [{"role": "system", "content": "ignore"}]}]:
            self.assertEqual(browser.post("/api/chat", json=data).status_code, 400)
        self.assertEqual(browser.post("/api/chat", json={"question": "x"}, headers={"Origin": "https://other.example"}).status_code, 403)

    def test_flask_ollama_failure_returns_actionable_error(self):
        self.seed()
        with patch.object(Ollama, "chat", side_effect=OllamaError("server unavailable")):
            browser = create_app(self.settings).test_client()
            result = browser.post("/api/chat", json={"question": "Avamar"})
            self.assertEqual(result.status_code, 502)
            self.assertEqual(result.json["error"], "server unavailable")

    def test_cli_load_list_and_failures(self):
        path = self.root / "notes.txt"
        path.write_text("<avamar>Restore files from backup</avamar>")
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(["load", "--db", self.settings.database, str(path)]), 0)
            self.assertEqual(main(["list", "--db", self.settings.database]), 0)
            self.assertEqual(main(["load", "--db", self.settings.database, str(self.root / "missing.txt")]), 1)

    def upload(self, browser, name, content, **extra):
        response = browser.post("/api/load", data={"files": (io.BytesIO(content), name), **extra})
        self.assertEqual(response.status_code, 200)
        return [json.loads(line) for line in response.data.splitlines()]

    def test_upload_tagged_notes_and_knowledge_inventory(self):
        self.seed()
        browser = create_app(self.settings).test_client()
        events = self.upload(browser, "Notes.txt", b"<vm-backup>Back up the VM using the approved backup product.</vm-backup>")
        self.assertEqual([e["type"] for e in events], ["loading", "loaded", "done"])
        self.assertEqual(events[1]["skills"], ["vm-backup"])
        self.assertEqual(StubOllama.calls, [])
        knowledge = browser.get("/api/knowledge").json
        self.assertEqual(knowledge["procedures"], 2)
        uploaded = next(file for file in knowledge["files"] if file["source"].startswith("upload://"))
        self.assertEqual(uploaded["filename"], "Notes.txt")
        self.assertEqual(uploaded["source"], "upload://Notes.txt")
        self.assertEqual(uploaded["skills"][0]["tag"], "vm-backup")
        self.assertTrue(self.store.search("VM backup"))
        result = browser.post("/api/chat", json={"question": "VM backup"})
        self.assertEqual(result.json["sources"][0]["tag"], "vm-backup")

    def test_upload_word_document_uses_filename(self):
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Restore the approved backup.")
        document.save(buffer)
        browser = create_app(self.settings).test_client()
        events = self.upload(browser, "Restore VM.docx", buffer.getvalue())
        self.assertEqual(events[1]["skills"], ["restore-vm"])
        self.assertEqual(StubOllama.calls, [])

    def test_upload_untagged_notes_infers_or_uses_offline_tags(self):
        browser = create_app(self.settings).test_client()
        events = self.upload(browser, "notes.txt", b"Check rack access")
        self.assertEqual(events[1]["skills"], ["rack-access-checklist"])
        self.assertEqual(len(StubOllama.calls), 1)
        StubOllama.calls.clear()
        events = self.upload(browser, "offline.txt", b"Check rack access", offline_tags="true")
        self.assertEqual(events[1]["skills"], ["check-rack-access"])
        self.assertEqual(StubOllama.calls, [])

    def test_reupload_replaces_entries_and_empty_import_preserves_them(self):
        browser = create_app(self.settings).test_client()
        self.upload(browser, "notes.txt", b"<avamar>Restore Avamar files</avamar>")
        self.upload(browser, "notes.txt", b"<vm-backup>Backup VM</vm-backup>")
        self.assertEqual(self.store.search("Avamar"), [])
        self.assertEqual(len(self.store.list_procedures()), 1)
        events = self.upload(browser, "notes.txt", b" ")
        self.assertEqual(events[1]["type"], "failed")
        self.assertEqual(events[-1]["failed"], 1)
        self.assertEqual(self.store.list_procedures()[0]["tag"], "vm-backup")

    def test_upload_batch_continues_after_bad_file_and_keeps_folder_identity(self):
        browser = create_app(self.settings).test_client()
        response = browser.post("/api/load", data={"files": [
            (io.BytesIO(b"bad document"), "broken.docx"),
            (io.BytesIO(b"<rack>Rack procedure</rack>"), "folder-a/notes.txt"),
            (io.BytesIO(b"<vm>VM procedure</vm>"), "folder-b/notes.txt"),
        ]})
        events = [json.loads(line) for line in response.data.splitlines()]
        self.assertEqual(events[-1], {"type": "done", "loaded": 2, "failed": 1})
        files = browser.get("/api/knowledge").json["files"]
        self.assertEqual([item["source"] for item in files], ["upload://folder-a/notes.txt", "upload://folder-b/notes.txt"])

    def test_failed_ai_label_preserves_previous_knowledge(self):
        browser = create_app(self.settings).test_client()
        self.upload(browser, "notes.txt", b"<rack>Rack procedure</rack>")
        with patch.object(Ollama, "chat", side_effect=OllamaError("cannot connect")):
            events = self.upload(browser, "notes.txt", b"Untagged replacement")
        self.assertEqual(events[1]["type"], "failed")
        self.assertIn("cannot connect", events[1]["error"])
        self.assertEqual(self.store.list_procedures()[0]["tag"], "rack")

    def test_upload_validation_and_path_traversal(self):
        browser = create_app(self.settings).test_client()
        self.assertEqual(browser.post("/api/load").status_code, 400)
        for name in ["../escape.txt", "/absolute.txt", "C:\\escape.txt", "notes.exe"]:
            events = self.upload(browser, name, b"<vm>Procedure</vm>")
            self.assertEqual(events[1]["type"], "failed", name)
        self.assertEqual(self.store.list_procedures(), [])
        self.assertFalse((self.root / "escape.txt").exists())
        result = browser.post("/api/load", data={"files": (io.BytesIO(b"text"), "notes.txt")}, headers={"Origin": "https://other.example"})
        self.assertEqual(result.status_code, 403)

    def test_upload_large_text_and_chat_request_limit(self):
        browser = create_app(self.settings).test_client()
        body = b"<large>" + b"Backup VM. " * 7000 + b"</large>"
        events = self.upload(browser, "large.txt", body)
        self.assertEqual(events[1]["type"], "loaded")
        result = browser.post("/api/chat", json={"question": "a" * 70000})
        self.assertEqual(result.status_code, 413)
        self.assertIn("64 KB", result.json["error"])

    def test_empty_knowledge_and_template_buttons(self):
        browser = create_app(self.settings).test_client()
        self.assertEqual(browser.get("/api/knowledge").json, {"files": [], "procedures": 0})
        page = browser.get("/").data.decode()
        self.assertIn('id="open-loader"', page)
        self.assertIn('id="show-knowledge"', page)
        self.assertIn("knowledge.js", page)
        self.assertIn("Active read timeout: 500 seconds", page)

    def test_explicit_timeout_overrides_environment_and_is_visible(self):
        with patch.dict("os.environ", {"OLLAMA_TIMEOUT": "120"}):
            from metropole.cli import build_parser
            args = build_parser().parse_args(["serve", "--timeout", "500"])
            self.assertEqual(args.timeout, 500)
            self.assertEqual(Settings.from_env().timeout, 120)
        settings = Settings(self.settings.host, database=self.settings.database, timeout=321)
        browser = create_app(settings).test_client()
        self.assertIn("Active read timeout: 321 seconds", browser.get("/").data.decode())
        self.assertEqual(browser.get("/api/status").json["timeout"], 321)

    def test_cancelled_upload_removes_temporary_copies(self):
        from tempfile import TemporaryDirectory as RealTemporaryDirectory
        staging_paths = []

        def staging(*args, **kwargs):
            directory = RealTemporaryDirectory(*args, **kwargs)
            staging_paths.append(Path(directory.name))
            return directory

        browser = create_app(self.settings).test_client()
        with patch("metropole.web.TemporaryDirectory", side_effect=staging):
            response = browser.post("/api/load", data={"files": (io.BytesIO(b"<vm>Backup VM</vm>"), "notes.txt")}, buffered=False)
            self.assertTrue(staging_paths[0].exists())
            response.close()
        self.assertFalse(staging_paths[0].exists())
        self.assertEqual(self.store.list_procedures(), [])

    def test_interactive_loader_and_chat(self):
        path = self.root / "notes.txt"
        path.write_text("<avamar>Confirm backup date before restoring Avamar files.</avamar>")
        opts = ["--db", self.settings.database, "--host", self.settings.host]
        with contextlib.redirect_stdout(io.StringIO()), patch("builtins.input", side_effect=[str(path), "/list", "/quit"]):
            self.assertEqual(main(["load", *opts]), 0)
        with contextlib.redirect_stdout(io.StringIO()) as output, patch("builtins.input", side_effect=["Restore Avamar files?", "Which date?", "/reset", "/quit"]):
            self.assertEqual(main(["chat", *opts]), 0)
        self.assertIn("[1]", output.getvalue())
        self.assertEqual(len(StubOllama.calls), 2)


if __name__ == "__main__":
    unittest.main()
